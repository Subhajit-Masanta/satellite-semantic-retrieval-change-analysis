"""Build the Project-I style report in the BCSE497J template format,
but for subject code BCSE407L (Computer Vision), in simple student English."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from pathlib import Path

ROOT = Path(r"C:\Users\subha\Desktop\projects\Computer Vision")
FIGS = ROOT / "report_figs"
OUT = ROOT / "BCSE407L_Project_Report_Satellite_Semantic_Retrieval_v3_draft.docx"
RFIG = ROOT / "results" / "figs"
import json, sys
sys.path.insert(0, str(ROOT / "report_tools"))
PG_FILE = ROOT / "report_tools" / "pages.json"
PAGES = json.loads(PG_FILE.read_text()) if PG_FILE.exists() else {}

doc = Document()

# ------------------------------------------------------------------ base style
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
st.paragraph_format.space_after = Pt(6)

for s in doc.sections:
    s.top_margin = Inches(1); s.bottom_margin = Inches(1)
    s.left_margin = Inches(1.25); s.right_margin = Inches(1)


def P(text="", size=12, bold=False, italic=False, align="left", spacing=1.15,
      space_after=6, upper=False, color=None):
    p = doc.add_paragraph()
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "just": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    p.paragraph_format.line_spacing = spacing
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text.upper() if upper else text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


def H1(text):            # 14, bold, upper, 1.5
    return P(text, size=14, bold=True, upper=True, spacing=1.5, space_after=8)


def H2(text):            # 13, bold, title case, 1.5
    return P(text, size=13, bold=True, spacing=1.5, space_after=6)


def H3(text):            # 12, bold italic, 1.5
    return P(text, size=12, bold=True, italic=True, spacing=1.5, space_after=4)


import sys as _sys
_sys.path.insert(0, str(ROOT / "report_tools"))
from noobify import noob


def body(text):
    text = noob(text)
    return P(text, size=12, align="just", spacing=1.15)


def bullet(text, bold_lead=None):
    text = noob(text)
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True
        r.font.name = "Times New Roman"; r.font.size = Pt(12)
    r = p.add_run(text)
    r.font.name = "Times New Roman"; r.font.size = Pt(12)
    return p


def numbered(text):
    text = noob(text)
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = "Times New Roman"; r.font.size = Pt(12)
    return p


def figure(path, caption, width=6.1):
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    P(caption, size=11, align="center", spacing=1.0, space_after=10)


def table(headers, rows, widths=None, font=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        pr = c.paragraphs[0]; pr.paragraph_format.space_after = Pt(2)
        r = pr.add_run(h); r.bold = True
        r.font.name = "Times New Roman"; r.font.size = Pt(font)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            pr = cells[i].paragraphs[0]; pr.paragraph_format.space_after = Pt(2)
            r = pr.add_run(str(v))
            r.font.name = "Times New Roman"; r.font.size = Pt(font)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return t


# ============================================================== TITLE PAGE
P("BCSE407L – Computer Vision", size=14, bold=True, align="center", spacing=1.5)
P("Project Report", size=13, bold=True, align="center", spacing=1.5, space_after=10)

P("SEMANTIC RETRIEVAL AND MULTI-TEMPORAL CHANGE ANALYSIS OF SATELLITE IMAGERY "
  "USING A SHARED FOUNDATION-MODEL EMBEDDING SPACE",
  size=16, bold=True, align="center", spacing=1.5, space_after=12)

P("Submitted by", size=12, align="center", spacing=1.5, space_after=8)
tbl = table(["Reg. No.", "Student Name"],
            [["23BCE0257", "AKANKSHA BHARTI"],
             ["23BCE0292", "SUBHAJIT MASANTA"],
             ["23BCE0318", "GONTLA VENKATA SIVA SAI NIKHIL"],
             ["23BCE2027", "ARJUN PRATAP SINGH"]],
            widths=[1.6, 3.4], font=12)
tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
P(space_after=4)

P("Under the Supervision of", size=12, align="center", spacing=1.5, space_after=6)
P("Dr. <Supervisor Name>", size=12, bold=True, align="center", spacing=1.5, space_after=0)
P("<Designation>", size=12, align="center", spacing=1.5, space_after=0)
P("School of Computer Science and Engineering", size=12, align="center",
  spacing=1.5, space_after=8)

P("In partial fulfilment of the requirements for the degree of", size=12,
  align="center", spacing=1.5, space_after=6)
P("B.Tech.", size=13, bold=True, align="center", spacing=1.5, space_after=0)
P("in", size=12, align="center", spacing=1.5, space_after=0)
P("Computer Science and Engineering", size=13, bold=True, align="center",
  spacing=1.5, space_after=8)
P("School of Computer Science and Engineering", size=12, align="center",
  spacing=1.5, space_after=12)
P("October 2026", size=12, bold=True, align="center", spacing=1.5)

doc.add_page_break()

# ================================================================= ABSTRACT
H1("Abstract")
import casual_ch1_4
_H = dict(P=P, H1=H1, H2=H2, H3=H3, body=body, bullet=bullet, numbered=numbered,
          table=table, figure=figure, doc=doc, FIGS=FIGS)
casual_ch1_4.abstract(_H)
P()
p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.15
r = p.add_run("Keywords - "); r.bold = True
r.font.name = "Times New Roman"; r.font.size = Pt(12)
r = p.add_run("Computer Vision, Remote Sensing, Satellite Imagery, Image "
              "Retrieval, Change Detection, CLIP, Embeddings, Vector Search.")
r.font.name = "Times New Roman"; r.font.size = Pt(12)

from docx.oxml import OxmlElement


def _page_field(footer):
    p = footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f = OxmlElement("w:fldSimple"); f.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r"); t_ = OxmlElement("w:t"); t_.text = "1"; r.append(t_); f.append(r)
    p._p.append(f)


def _numfmt(section, fmt, start):
    sp = section._sectPr
    for old in sp.findall(qn("w:pgNumType")):
        sp.remove(old)
    n = OxmlElement("w:pgNumType"); n.set(qn("w:fmt"), fmt); n.set(qn("w:start"), str(start))
    sp.find(qn("w:pgMar")).addnext(n)          # schema order: right after pgMar


front = doc.sections[0]
front.different_first_page_header_footer = True      # no number on the title page
_page_field(front.footer); _numfmt(front, "lowerRoman", 0)
main = doc.add_section(WD_SECTION.NEW_PAGE)
main.different_first_page_header_footer = False
main.footer.is_linked_to_previous = False
_page_field(main.footer); _numfmt(main, "decimal", 1)

# ===================================================== CHAPTERS 1-4 (casual)
casual_ch1_4.add(_H)

import new_sections
new_sections.add(dict(P=P, H1=H1, H2=H2, H3=H3, body=body, bullet=bullet,
                      numbered=numbered, table=table, doc=doc, RFIG=RFIG,
                      ROOT=ROOT, Pt=Pt, Inches=Inches, qn=qn,
                      WD_ALIGN_PARAGRAPH=WD_ALIGN_PARAGRAPH))

doc.save(OUT)
print("saved:", OUT)
